import copy
import io

import docx
import pymupdf
import pytest

from backend.resume.apply import ProtectedFactError, build_tailored_profile
from backend.resume.docx_generator import build_resume_docx
from backend.resume.pdf_generator import build_resume_pdf

CSX0 = "Built real-time rail operations dashboards in Angular used by 300+ dispatchers."


def session_with(statuses: dict, edited: dict | None = None) -> dict:
    changes = [
        {"id": "a", "targetBulletId": "exp-0-b0", "original": CSX0, "updated": "ACCEPTED: dashboards in Angular and TypeScript."},
        {"id": "r", "targetBulletId": "exp-0-b2", "original": "Deployed frontend services to AWS with automated CI pipelines.",
         "updated": "REJECTED line"},
        {"id": "p", "targetBulletId": "exp-1-b1", "original": "Implemented React components for case-worker workflows.",
         "updated": "PENDING line"},
        {"id": "e", "targetBulletId": "summary",
         "original": "Frontend engineer with 5 years of experience building enterprise web applications with Angular and TypeScript. Focused on accessible, high-performance UIs.",
         "updated": "proposed summary", "editedText": (edited or {}).get("e")},
    ]
    for change in changes:
        change["status"] = statuses.get(change["id"], "pending")
    return {"jobTitle": "Frontend Engineer", "changes": changes}


def pdf_text(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as document:
        return "\n".join(page.get_text() for page in document)


def docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)


def test_only_accepted_and_edited_changes_applied(profile):
    master = copy.deepcopy(profile)
    session = session_with({"a": "accepted", "r": "rejected", "p": "pending", "e": "edited"}, {"e": "EDITED summary text."})
    tailored, applied, skipped = build_tailored_profile(profile, session)
    assert applied == ["e", "a"] and skipped == []
    assert tailored["summary"]["text"] == "EDITED summary text."
    bullets = {b["id"]: b["text"] for e in tailored["experience"] for b in e["bullets"]}
    assert bullets["exp-0-b0"].startswith("ACCEPTED")
    assert bullets["exp-0-b2"] == "Deployed frontend services to AWS with automated CI pipelines."
    assert bullets["exp-1-b1"] == "Implemented React components for case-worker workflows."
    assert profile == master  # Master Resume never mutated


def test_stale_change_is_skipped(profile):
    session = session_with({"a": "accepted"})
    session["changes"][0]["original"] = "An older version of the bullet."
    tailored, applied, skipped = build_tailored_profile(profile, session)
    assert applied == [] and "original line changed" in skipped[0]


def test_protected_facts_guard(profile, monkeypatch):
    from backend.resume import apply as apply_module
    real = apply_module.protected_facts
    calls = iter([["tampered"], real(profile)])
    monkeypatch.setattr(apply_module, "protected_facts", lambda p: next(calls))
    with pytest.raises(ProtectedFactError):
        build_tailored_profile(profile, None)


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_generated_files_contain_approved_text_only(profile, fmt):
    tailored, _, _ = build_tailored_profile(profile, session_with({"a": "accepted", "r": "rejected", "p": "pending"}))
    data = build_resume_pdf(tailored) if fmt == "pdf" else build_resume_docx(tailored)
    text = pdf_text(data) if fmt == "pdf" else docx_text(data)
    assert data[:4] == (b"%PDF" if fmt == "pdf" else b"PK\x03\x04")
    for expected in ["Alex Rivera", "Summary", "Skills", "Experience", "Software Engineer — CSX", "Education",
                     "B.S. Computer Science, University of Florida (2019)", "Certifications",
                     "AWS Certified Cloud Practitioner — Amazon Web Services — 2022", "ACCEPTED: dashboards in Angular and TypeScript."]:
        assert expected in text, expected
    assert "REJECTED" not in text and "PENDING" not in text


def test_docx_is_ats_friendly(profile):
    document = docx.Document(io.BytesIO(build_resume_docx(profile)))
    assert len(document.tables) == 0 and len(document.inline_shapes) == 0
    styles = [p.style.name for p in document.paragraphs]
    assert styles.count("Heading 2") == 5 and styles.count("List Bullet") == 5


def test_pdf_escapes_markup_and_handles_empty_profile():
    text = pdf_text(build_resume_pdf({"personalInformation": {"fullName": "A & B <Co>"}, "summary": {"text": "x < y & z"}}))
    assert "A & B <Co>" in text and "x < y & z" in text
    assert pdf_text(build_resume_pdf({})) == ""
    assert docx.Document(io.BytesIO(build_resume_docx({})))
