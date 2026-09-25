"""Resume file -> raw text. Nothing else (no AI, no OCR).

Port of resume-tailor/server/resumeParser.js with PyMuPDF (PDF) and
python-docx (DOCX), plus stricter validation: extension, size, file
signature, container integrity, empty text.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Optional

# Vercel Functions reject request bodies over 4.5 MB before our code runs, so
# the app limit is 4 MB (leaves room for multipart overhead) everywhere —
# the same limit locally and in production. (Standalone app: 10 MB.)
MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_FILE_LABEL = "4 MB"

RESUME_TYPES = ("pdf", "docx")
DOCUMENT_TYPES = ("pdf", "docx", "txt")  # job descriptions may also be .txt

ACCEPTED_MIME = {
    "pdf": {"application/pdf", "application/x-pdf"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    "txt": {"text/plain"},
}
# Browsers/OSes often send these for any file; don't reject on them.
GENERIC_MIME = {"", "application/octet-stream", "binary/octet-stream", "application/zip", "application/x-zip-compressed"}


class ResumeParseError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass
class ParsedDocument:
    text: str
    file_type: str
    filename: str
    size_bytes: int
    sha256: str
    content_type: Optional[str] = None
    page_count: Optional[int] = None

    def metadata(self) -> dict:
        return {
            "filename": self.filename,
            "contentType": self.content_type,
            "sizeBytes": self.size_bytes,
            "sha256": self.sha256,
            "fileType": self.file_type,
            "pageCount": self.page_count,
            "textLength": len(self.text),
            "uploadedAt": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
        }


def _file_type(filename: str, allowed: tuple[str, ...]) -> str:
    name = (filename or "").lower()
    for kind in allowed:
        if name.endswith("." + kind):
            return kind
    raise ResumeParseError(
        "unsupported_type",
        f"Unsupported file type. Please upload a {' or '.join('.' + k for k in allowed)} file.",
    )


def _check_signature(data: bytes, filename: str, kind: str) -> None:
    if kind == "pdf" and not data[:8].startswith(b"%PDF-"):
        raise ResumeParseError(
            "bad_signature",
            f'"{filename}" doesn\'t look like a real PDF (its content doesn\'t start with the PDF file signature). '
            "It may be an HTML export or another format renamed to .pdf — try re-exporting/re-saving it as an actual PDF.",
        )
    if kind == "docx" and data[:2] != b"PK":
        raise ResumeParseError(
            "bad_signature",
            f'"{filename}" doesn\'t look like a real .docx file (its content doesn\'t start with the ZIP file signature '
            ".docx files use). Try re-saving it from Word/Google Docs as .docx.",
        )


def _check_mime(content_type: Optional[str], kind: str, filename: str) -> None:
    mime = (content_type or "").split(";")[0].strip().lower()
    if mime in GENERIC_MIME or mime in ACCEPTED_MIME[kind]:
        return
    raise ResumeParseError("bad_mime", f'"{filename}" was uploaded as {mime}, which is not a .{kind} file.')


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def parse_pdf(data: bytes) -> tuple[str, int]:
    import pymupdf

    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        raise ResumeParseError("malformed", "This PDF could not be opened — it may be corrupted.") from None
    with document:
        if document.needs_pass:
            raise ResumeParseError("encrypted", "This PDF is password-protected. Remove the password and try again.")
        try:
            pages = [page.get_text("text") for page in document]
        except Exception:
            raise ResumeParseError("malformed", "This PDF could not be read — it may be corrupted.") from None
        return _tidy("\n".join(pages)), document.page_count


def parse_docx(data: bytes) -> str:
    import docx
    from docx.table import Table

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ResumeParseError("malformed", "This .docx file has no document body — try re-saving it from Word/Google Docs.")
        document = docx.Document(io.BytesIO(data))
    except ResumeParseError:
        raise
    except Exception:
        raise ResumeParseError("malformed", "This .docx file could not be opened — it may be corrupted.") from None

    lines: list[str] = []
    for block in document.iter_inner_content():  # paragraphs and tables, in document order
        if isinstance(block, Table):
            seen = set()
            for row in block.rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    lines.extend(p.text for p in cell.paragraphs)
        else:
            lines.append(block.text)
    return _tidy("\n".join(lines))


def parse_document(
    filename: str, data: bytes, content_type: Optional[str] = None, *, allowed: tuple[str, ...] = RESUME_TYPES
) -> ParsedDocument:
    kind = _file_type(filename, allowed)
    if not data:
        raise ResumeParseError("empty_file", f'"{filename}" is empty.')
    if len(data) > MAX_FILE_BYTES:
        raise ResumeParseError("too_large", f"File is too large (max {MAX_FILE_LABEL}).")
    _check_mime(content_type, kind, filename)

    page_count = None
    if kind == "txt":
        try:
            text = _tidy(data.decode("utf-8-sig"))
        except UnicodeDecodeError:
            raise ResumeParseError("malformed", "This text file is not valid UTF-8.") from None
    else:
        _check_signature(data, filename, kind)
        if kind == "pdf":
            text, page_count = parse_pdf(data)
        else:
            text = parse_docx(data)

    if not text.strip():
        raise ResumeParseError(
            "no_text",
            "Could not extract any text from this file. If it is a scanned/image-based PDF, try a text-based PDF or "
            "DOCX instead (OCR is not supported).",
        )
    return ParsedDocument(
        text=text, file_type=kind, filename=filename, size_bytes=len(data),
        sha256=hashlib.sha256(data).hexdigest(), content_type=content_type, page_count=page_count,
    )


def parse_resume(filename: str, data: bytes, content_type: Optional[str] = None) -> ParsedDocument:
    return parse_document(filename, data, content_type, allowed=RESUME_TYPES)
