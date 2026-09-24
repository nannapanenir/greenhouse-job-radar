"""Tailored profile -> ATS-friendly .docx (port of docxGenerator.js).

Plain paragraphs and headings only: no tables, text boxes, columns or images.
"""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt


def _contact_line(personal: dict) -> str:
    return "  |  ".join(v for v in (personal.get("email"), personal.get("phone"), personal.get("location"), personal.get("linkedIn")) if v)


def _spacing(paragraph, before: int = 0, after: int = 0) -> None:
    # docx library "twips" in the JS version (240 = 12pt); python-docx uses points.
    paragraph.paragraph_format.space_before = Pt(before / 20)
    paragraph.paragraph_format.space_after = Pt(after / 20)


def build_resume_docx(profile: dict) -> bytes:
    personal = profile.get("personalInformation") or {}
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def heading(text: str) -> None:
        _spacing(document.add_heading(text, level=2), before=240, after=120)

    name = document.add_paragraph()
    run = name.add_run(personal.get("fullName") or "")
    run.bold = True
    run.font.size = Pt(16)
    _spacing(name, after=80)

    contact = _contact_line(personal)
    if contact:
        paragraph = document.add_paragraph()
        paragraph.add_run(contact).font.size = Pt(10)
        _spacing(paragraph, after=200)

    summary = (profile.get("summary") or {}).get("text")
    if summary:
        heading("Summary")
        _spacing(document.add_paragraph(summary), after=120)

    skills = profile.get("skills") or []
    if skills:
        heading("Skills")
        _spacing(document.add_paragraph(", ".join(s.get("name", "") for s in skills)), after=120)

    experience = profile.get("experience") or []
    if experience:
        heading("Experience")
        for exp in experience:
            title = document.add_paragraph()
            title.add_run(f"{exp.get('title') or ''} — {exp.get('company') or ''}").bold = True
            _spacing(title, before=160, after=20)
            date_line = "  |  ".join(v for v in (f"{exp.get('startDate') or ''} – {exp.get('endDate') or ''}", exp.get("location")) if v)
            if date_line.strip():
                paragraph = document.add_paragraph()
                run = paragraph.add_run(date_line)
                run.italic = True
                run.font.size = Pt(10)
                _spacing(paragraph, after=60)
            for bullet in exp.get("bullets") or []:
                _spacing(document.add_paragraph(bullet.get("text", ""), style="List Bullet"), after=40)

    education = profile.get("education") or []
    if education:
        heading("Education")
        for edu in education:
            end = f" ({edu['endDate']})" if edu.get("endDate") else ""
            _spacing(document.add_paragraph(f"{edu.get('degree') or ''}, {edu.get('institution') or ''}{end}"), after=40)

    certifications = profile.get("certifications") or []
    if certifications:
        heading("Certifications")
        for cert in certifications:
            line = " — ".join(v for v in (cert.get("name"), cert.get("issuer"), cert.get("year")) if v)
            _spacing(document.add_paragraph(line), after=40)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
